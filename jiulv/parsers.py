"""文档解析：把各种格式统一成纯文本。

docx / pdf / xlsx 属于可选依赖，没装就跳过并在结果里说明，不影响主流程。
"""

from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple

TEXT_EXTS = {".md", ".markdown", ".txt", ".csv", ".json", ".yaml", ".yml"}
SUPPORTED_EXTS = TEXT_EXTS | {".docx", ".pdf", ".xlsx"}


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_EXTS


def read_text_file(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return path.read_bytes().decode("utf-8", errors="ignore")


def _docx_text(path: Path) -> str:
    try:
        import docx  # type: ignore
    except ImportError:
        raise RuntimeError("需要 python-docx：pip install 'jiulv-kb[parsers]'")
    document = docx.Document(str(path))
    parts: List[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        if style.startswith("heading") or style.startswith("标题"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            parts.append("#" * min(int(level), 6) + " " + text)
        else:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        raise RuntimeError("需要 pypdf：pip install 'jiulv-kb[parsers]'")
    reader = PdfReader(str(path))
    out: List[str] = []
    for page in reader.pages:
        try:
            out.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n\n".join(p.strip() for p in out if p.strip())


def _xlsx_text(path: Path) -> str:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        raise RuntimeError("需要 openpyxl：pip install 'jiulv-kb[parsers]'")
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out: List[str] = []
    for sheet in wb.worksheets:
        out.append(f"## {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c not in (None, "")]
            if cells:
                out.append(" | ".join(cells))
    wb.close()
    return "\n".join(out)


def _legacy_office_text(path: Path) -> Optional[str]:
    """老格式 .doc/.xls 不做解析，但如果是改了后缀的 zip 就尽量捞一层。"""
    try:
        with zipfile.ZipFile(path) as zf:
            for name in zf.namelist():
                if name.endswith(".xml") and "document" in name:
                    return zf.read(name).decode("utf-8", errors="ignore")
    except Exception:
        return None
    return None


def extract_text(path: Path) -> Tuple[str, str]:
    """返回 (文本, 格式说明)。失败时抛 RuntimeError，由调用方决定跳过还是中止。"""
    ext = path.suffix.lower()
    if ext in TEXT_EXTS:
        return read_text_file(path), "text"
    if ext == ".docx":
        return _docx_text(path), "docx"
    if ext == ".pdf":
        return _pdf_text(path), "pdf"
    if ext == ".xlsx":
        return _xlsx_text(path), "xlsx"
    fallback = _legacy_office_text(path)
    if fallback:
        return fallback, "zip-xml"
    raise RuntimeError(f"不支持的格式：{ext}")


def parser_status() -> Dict[str, bool]:
    """检查可选依赖是否就绪，供 CLI 提示。"""
    status = {}
    for mod, label in (("docx", "python-docx"), ("pypdf", "pypdf"), ("openpyxl", "openpyxl")):
        try:
            __import__(mod)
            status[label] = True
        except ImportError:
            status[label] = False
    return status
